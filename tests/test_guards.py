"""방어선 회귀 — 내보내기·경로·수집 복원력·서버 계약.

여기 있는 것들은 **자매 저장소가 이미 값을 치른 실패**다. 새로 짓는 김에
처음부터 막으려고 옮겨 심었고, 각 테스트는 그 사고를 그대로 재현한다.
"""
from __future__ import annotations

import csv
import json
import sqlite3

import pytest

from law_mcp import exporters
from law_mcp.client import LawClient, LawError
from law_mcp.config import TARGETS, scrub
from law_mcp.exporters import export, extra_columns, safe_name
from law_mcp.models import COLUMNS, Record


def _recs():
    return [
        Record(doc_id="1", target="law", doc_type="현행법령", title="초·중등교육법",
               number="00311", promulgated="2019-07-15", authority="교육부",
               raw={"법령일련번호": "1", "법령ID": "008770", "소관부처코드": "1342000",
                    "공동부령정보": ""}),
        Record(doc_id="2", target="prec", doc_type="판례", title="교육세부과처분취소",
               number="2025두35074", promulgated="2026-02-12",
               raw={"판례일련번호": "2", "데이터출처명": "대법원", "선고": "선고"}),
    ]


# ── 🔴 미매핑 필드가 파일에 나오는가 ────────────────────────────────────────
# 자매 저장소는 정규화 표에 없는 필드를 MCP 응답·csv·xlsx 어디에도 내보내지 않아
# 국회의안의 '제안이유 및 주요내용'(97% 채워진 서술형 본문)을 통째로 잃었다.
def test_unmapped_fields_become_columns():
    cols = extra_columns(_recs())
    assert "법령ID" in cols and "데이터출처명" in cols
    # 값이 전부 빈 필드는 열로 만들지 않는다(표만 넓어진다)
    assert "공동부령정보" not in cols


def test_csv_contains_unmapped_values(tmp_path):
    export(_recs(), ["csv"], str(tmp_path), "t")
    rows = list(csv.DictReader(open(tmp_path / "t.csv", encoding="utf-8-sig")))
    assert rows[0]["법령ID"] == "008770"
    assert rows[1]["데이터출처명"] == "대법원"


def test_xlsx_contains_unmapped_values(tmp_path):
    from openpyxl import load_workbook
    export(_recs(), ["xlsx"], str(tmp_path), "t")
    ws = load_workbook(tmp_path / "t.xlsx").active
    header = [c.value for c in ws[1]]
    assert "법령ID" in header and "데이터출처명" in header
    assert header[:len(COLUMNS)] == COLUMNS, "정규화 열이 앞에 와야 한다"


def test_sqlite_and_json_carry_raw(tmp_path):
    export(_recs(), ["sqlite", "json"], str(tmp_path), "t")
    con = sqlite3.connect(tmp_path / "t.sqlite")
    raw = con.execute("SELECT raw FROM records LIMIT 1").fetchone()[0]
    con.close()
    assert json.loads(raw)["법령ID"] == "008770"
    data = json.loads((tmp_path / "t.json").read_text(encoding="utf-8"))
    assert data[0]["citation"]["발령주체"] == "교육부"


def test_xlsx_oversized_cell_is_truncated_not_fatal(tmp_path):
    """엑셀 셀 상한(32,767자)을 넘으면 openpyxl 이 죽어 **파일 전체가 안 써진다**.

    조문 본문은 그 길이를 쉽게 넘는다 — 잘라서라도 저장하고 잘렸다고 표시한다.
    """
    r = Record(doc_id="1", target="law", title="x", raw={"조문내용": "가" * 40_000})
    export([r], ["xlsx"], str(tmp_path), "big")
    from openpyxl import load_workbook
    ws = load_workbook(tmp_path / "big.xlsx").active
    cell = [c.value for c in ws[2]][-1]
    assert len(cell) <= 32_767 and "잘림" in cell


# ── 경로·형식 ──────────────────────────────────────────────────────────────
@pytest.mark.parametrize("bad", ["../escaped", "..\\escaped", "con", "  ", "a/b"])
def test_safe_name_cannot_escape(bad):
    s = safe_name(bad)
    assert "/" not in s and "\\" not in s and ".." not in s
    assert s.upper().split(".")[0] not in {"CON", "PRN", "AUX", "NUL"}


def test_export_validates_all_formats_before_writing_anything(tmp_path):
    """['json','bogus'] 가 json 을 쓴 뒤 죽으면 메타는 사라지고 쿼터는 이미 태운 뒤다."""
    with pytest.raises(ValueError):
        export(_recs(), ["json", "bogus"], str(tmp_path), "t")
    assert list(tmp_path.iterdir()) == [], "아무 파일도 쓰지 않아야 한다"


def test_export_accepts_bare_string_format(tmp_path):
    """문자열을 넘기면 문자 단위로 순회해 '지원하지 않는 형식: j' 가 났었다."""
    paths = export(_recs(), "json", str(tmp_path), "t")
    assert len(paths) == 1 and paths[0].endswith(".json")


# ── 🔴 한 축이 실패해도 나머지를 버리지 않는가 ─────────────────────────────
class _StubClient(LawClient):
    """`search` 만 갈아끼운다 — 네트워크 없이 수집 복원력을 시험한다."""

    def __init__(self, failing: set[tuple[str, str]]):
        super().__init__(oc="stub", throttle=0.0)
        self.failing = failing
        self.seen: list[tuple[str, str]] = []

    def search(self, target, query, **kw):   # type: ignore[override]
        self.seen.append((target, query))
        if (target, query) in self.failing:
            return [], {"target": target, "target_label": target, "query": query,
                        "total": 0, "fetched": 0, "error": "일시 오류(모의)"}
        r = Record(doc_id=f"{target}-{query}", target=target, title=f"{target}/{query}")
        return [r], {"target": target, "target_label": target, "query": query,
                     "total": 1, "fetched": 1}


def test_one_failing_axis_does_not_discard_the_rest():
    """자매 저장소는 첫 페이지 실패를 그대로 올려 **이미 모은 것 전량**을 잃었다."""
    c = _StubClient(failing={("law", "평생교육")})
    recs, meta = c.collect(["law", "prec"], ["교육", "평생교육"], max_records=100)
    assert len(recs) == 3, "실패한 축 하나만 빠지고 나머지는 남아야 한다"
    assert meta["failed_axes"] == ["law/평생교육"]
    assert "failed_axes_note" in meta
    assert len(c.seen) == 4, "실패 뒤에도 남은 축을 계속 조회해야 한다"


def test_budget_exhaustion_is_reported_not_silent():
    """max_records 는 전체 예산이다 — 뒤 축이 조회조차 안 된 사실을 알려야 한다."""
    c = _StubClient(failing=set())
    recs, meta = c.collect(["law", "prec", "ordin"], ["교육"], max_records=2)
    assert meta["axes_unsearched"], "조회되지 않은 축이 보고돼야 한다"
    assert "stopped_early_note" in meta
    assert meta["axes_planned"] == 3 and meta["axes_searched"] == 2


def test_collect_validates_every_target_before_starting():
    c = _StubClient(failing=set())
    with pytest.raises(ValueError):
        c.collect(["law", "zzz"], ["교육"], max_records=10)
    assert c.seen == [], "하나라도 틀리면 시작하지 않는다"


# ── 클라이언트 계약 ────────────────────────────────────────────────────────
def test_missing_oc_fails_before_any_request():
    with pytest.raises(LawError) as e:
        LawClient(oc="").search("law", "교육")
    assert "LAW_OC" in str(e.value)


def test_extra_params_cannot_override_validated_ones():
    c = LawClient(oc="stub", throttle=0.0)
    with pytest.raises(LawError) as e:
        c.search("law", "교육", extra={"query": "*"})
    assert "예약 파라미터" in str(e.value)


def test_body_requires_identifier():
    with pytest.raises(LawError) as e:
        LawClient(oc="stub", throttle=0.0).body("law", "  ")
    assert "MST" in str(e.value), "갈래별 파라미터명을 알려 줘야 한다"


# ── 누출 ───────────────────────────────────────────────────────────────────
def test_scrub_removes_secretish_params_but_keeps_OC():
    """OC 는 비밀값이 아니고 상세링크의 정상적인 일부다 — 지우면 링크가 깨진다."""
    s = scrub("https://x/?OC=someone&serviceKey=abc123%2F&type=XML")
    assert "serviceKey=***" in s
    assert "OC=someone" in s


# ── 서버 계약 ──────────────────────────────────────────────────────────────
def test_stdio_is_the_default_transport():
    """기본 전송이 바뀌면 기존 MCP 등록이 전부 죽는다."""
    from law_mcp import server
    p = server.argparse.ArgumentParser(prog="law-mcp")
    p.add_argument("--transport", default="stdio")
    assert p.parse_args([]).transport == "stdio"


def test_tools_return_dict_not_exception(monkeypatch):
    """도구는 어떤 예외도 밖으로 흘리지 않는다 — 항상 JSON 직렬화 가능한 dict."""
    from law_mcp import server
    monkeypatch.setattr(server, "get_oc", lambda: "stub")

    fn = server.law_search.fn if hasattr(server.law_search, "fn") else server.law_search
    out = fn(target="zzz", query="교육")
    assert isinstance(out, dict) and "error" in out
    json.dumps(out, ensure_ascii=False)


def test_law_targets_groups_instead_of_dumping_72_rows():
    """72종을 평평하게 늘어놓으면 호출자가 고를 수 없고 응답만 길어진다."""
    from law_mcp import server
    fn = server.law_targets.fn if hasattr(server.law_targets, "fn") else server.law_targets
    out = fn()
    assert out["갈래_수"] == 72
    groups = out["갈래"]
    assert "법령해석(부처 39종)" in groups
    assert len(groups["법령해석(부처 39종)"]) == 39
    assert "law" in groups["법령·규칙"]
    # 각 갈래는 정확히 한 묶음에만 들어가야 한다(중복·누락 없이 72종)
    flat = [c for codes in groups.values() for c in codes]
    assert len(flat) == 72 and len(set(flat)) == 72
    assert set(out["본문_없는_갈래"]) == {"licbyl", "ordinbyl", "moefCgmExpc", "ntsCgmExpc"}
    assert "lsHistory" in out["다루지_않는_target"]
    json.dumps(out, ensure_ascii=False)


def test_unknown_target_error_does_not_dump_every_target():
    """오류 메시지가 72종을 전부 뱉으면 화면을 덮는다 — 묶음 요약이어야 한다."""
    from law_mcp.config import resolve_target
    with pytest.raises(ValueError) as e:
        resolve_target("zzz")
    msg = str(e.value)
    assert "law_targets" in msg
    assert msg.count("CgmExpc") <= 6, "부처 39종을 전부 나열하고 있다"


def test_registry_entries_are_internally_consistent():
    """레지스트리가 복사·붙여넣기로 망가졌는지 보는 최소 검사.

    ⚠️ `id_param` 을 `("MST","ID")` 로 못박고 있었는데 **틀렸다** — `lstrm` 은
       `trmSeqs` 다(실측). 내가 쓴 테스트가 내 짐작을 굳히고 있었다.
    """
    for code, t in TARGETS.items():
        assert t.code == code, code
        assert t.id_param, code
        assert t.id_field, code
        assert t.list_root and t.record_tag, code
        assert t.label, code
        # 본문이 있다면 루트 태그를 알아야 하고, 없다면 비어 있어야 한다.
        assert bool(t.body_root) == t.has_body, code


def test_registry_covers_the_measured_catalog():
    """2026-09-08 전수 실측(62종) + 초판 10종 = 72종."""
    assert len(TARGETS) == 72
    # T41 이 미해결로 남긴 부처 1차 법령해석 — 39종 전부 들어와야 한다.
    ministry = [c for c in TARGETS if c.endswith("CgmExpc")]
    assert len(ministry) == 39
    assert "moeCgmExpc" in TARGETS, "교육부 법령해석"


def test_ministry_targets_share_one_shape():
    """부처 39종은 실측상 전부 같은 모양이다 — 하나라도 어긋나면 표가 오염된 것이다."""
    for code, t in TARGETS.items():
        if not code.endswith("CgmExpc"):
            continue
        assert t.list_root == "CgmExpc", code
        assert t.record_tag == "cgmExpc", code
        assert t.id_param == "ID", code
        assert t.id_field == "법령해석일련번호", code


def test_lstrm_uses_its_own_identifier_param():
    """🔴 `MST`/`ID` 만 있는 줄 알았다 — 법령용어는 `trmSeqs` 다(상세링크가 알려줬다)."""
    assert TARGETS["lstrm"].id_param == "trmSeqs"
    assert TARGETS["lstrm"].id_field == "법령용어ID"


def test_record_tag_mismatch_is_the_norm_not_the_exception():
    """45/62 가 제 이름이 아닌 레코드 태그를 쓴다 — 규칙으로 유추하면 대부분 틀린다."""
    mismatched = [c for c, t in TARGETS.items() if t.record_tag.lower() != c.lower()]
    assert len(mismatched) >= 45
    assert "school" in mismatched and TARGETS["school"].record_tag == "admrul"
    assert TARGETS["ttSpecialDecc"].record_tag == "decc"
    assert TARGETS["admbyl"].record_tag == "admrulbyl"


def test_targets_without_xml_body_are_marked():
    """별표·서식 본문과 일부 부처 해석은 XML 본문이 없다(HTML). 표에 적어 둔다."""
    no_body = {c for c, t in TARGETS.items() if not t.has_body}
    assert no_body == {"licbyl", "ordinbyl", "moefCgmExpc", "ntsCgmExpc"}


def test_body_refused_for_targets_without_xml_body():
    """본문이 없는 갈래는 **호출 전에** 막는다 — 부르면 HTML 을 받고 파서가 헛돈다."""
    with pytest.raises(LawError) as e:
        LawClient(oc="stub", throttle=0.0).body("licbyl", "12345")
    assert "본문" in str(e.value)
    assert "별표서식파일링크" in str(e.value), "대신 무엇을 보라고 알려 줘야 한다"


def test_log_scrubber_does_not_break_numeric_format_args(caplog):
    """🔴 초판은 모든 로그 인자를 str() 로 바꿔 `%d` 포맷을 터뜨렸다.

    `log.warning("재시도 %d/%d", 1, 3)` 이 TypeError 를 냈다 — 실제로 쓰이는
    경고 경로였고, 회귀 테스트가 아니었으면 라이브에서야 드러났을 것이다.
    """
    import logging as _logging
    from law_mcp.config import install_log_scrubber

    install_log_scrubber()
    log = _logging.getLogger("law_mcp")
    with caplog.at_level(_logging.WARNING, logger="law_mcp"):
        log.warning("재시도 %d/%d — key=%s", 1, 3, "serviceKey=abc")
    assert "재시도 1/3" in caplog.text


# ── display 절삭 오탐 ───────────────────────────────────────────────────────
class _EchoClient(LawClient):
    """`numOfRows` 에코가 붙은 응답을 흉내낸다(네트워크 없음)."""

    def __init__(self, total: int, echo: int, returned: int):
        super().__init__(oc="stub", throttle=0.0)
        self._t, self._e, self._r = total, echo, returned

    def _page(self, tgt, params, *, expect_records):   # type: ignore[override]
        recs = [Record(doc_id=str(i), target=tgt.code, title=f"t{i}")
                for i in range(self._r)]
        return self._t, recs, {"numOfRows": str(self._e)}


def test_short_result_is_not_mistaken_for_a_capped_page_size():
    """🔴 `numOfRows` 는 요청 크기가 아니라 **실제 반환 건수**다.

    실측: display=100 · total=3 → numOfRows=3. `echo != size` 만 보면 결과가
    적을 때마다 거짓 경고가 뜨고, size 가 3으로 줄어 이후 페이징이 기어간다.
    """
    _, meta = _EchoClient(total=3, echo=3, returned=3).search(
        "law", "초중등교육", max_records=50, page_size=100)
    assert "page_size_note" not in meta
    assert meta["page_size"] == 100


def test_page_size_is_clamped_before_the_request():
    """실측 상한은 500이다. 1000을 넘겨도 **우리가 먼저** 500으로 줄여 보낸다.

    그래서 '서버가 1000을 500으로 깎았다'는 상황은 정상 경로에서는 생기지 않는다.
    """
    _, meta = _EchoClient(total=8517, echo=500, returned=500).search(
        "ordin", "교육", max_records=500, page_size=1000)
    assert meta["page_size"] == 500
    assert "page_size_note" not in meta


def test_server_lowering_its_limit_is_reported():
    """상한이 바뀌면(예: 500→300) 조용히 적게 받는다 — 그때는 알려야 한다.

    지금은 일어나지 않지만, 이 경고가 없으면 상한이 내려간 날 수집이 조용히 준다.
    """
    _, meta = _EchoClient(total=8517, echo=300, returned=300).search(
        "ordin", "교육", max_records=500, page_size=500)
    assert "page_size_note" in meta and "깎였습니다" in meta["page_size_note"]


# ── 배포 메타 ───────────────────────────────────────────────────────────────
def test_registry_description_within_100_chars():
    """🔴 MCP 레지스트리는 `description` 을 100자로 제한한다.

    자매 저장소는 107자로 넘겨 **바이너리 3종을 다 빌드하고 Release 까지 만든 뒤**
    마지막 발행 단계에서 422 로 거부됐다. 제약을 여기서 붙든다.
    """
    import json
    from pathlib import Path
    root = Path(__file__).resolve().parents[1]
    desc = json.loads((root / "server.json").read_text(encoding="utf-8"))["description"]
    assert len(desc) <= 100, f"{len(desc)}자 — 레지스트리가 100자에서 거부한다"


def test_packaging_versions_agree():
    """pyproject·server.json·mcpb 의 버전이 어긋나면 릴리스가 조용히 뒤섞인다."""
    import json
    from pathlib import Path
    root = Path(__file__).resolve().parents[1]
    pyproject = (root / "pyproject.toml").read_text(encoding="utf-8")
    ver = next(l.split("=")[1].strip().strip('"')
               for l in pyproject.splitlines() if l.startswith("version ="))
    assert json.loads((root / "server.json").read_text(encoding="utf-8"))["version"] == ver
    assert json.loads((root / "mcpb" / "manifest.json").read_text(encoding="utf-8"))["version"] == ver


def test_mcpb_manifest_lists_every_registered_tool():
    """도구를 추가하고 매니페스트를 안 고치면 배포본에서만 목록이 낡는다."""
    import asyncio
    import json
    from pathlib import Path
    from law_mcp import server
    root = Path(__file__).resolve().parents[1]
    listed = {t["name"] for t in json.loads(
        (root / "mcpb" / "manifest.json").read_text(encoding="utf-8"))["tools"]}
    # 권위는 FastMCP 의 등록부다. 이 버전의 `@mcp.tool()` 은 함수를 그대로 돌려주므로
    # 모듈 속성만 훑으면 데코레이터 구현이 바뀔 때 조용히 빈 집합이 된다.
    actual = {t.name for t in asyncio.run(server.mcp.list_tools())}
    assert actual, "도구를 하나도 못 찾았다 — 탐지 방법이 낡았다"
    assert actual == listed, f"매니페스트와 어긋남: {actual ^ listed}"


def test_eflaw_body_refuses_without_ef_date():
    """`efYd` 없이 부르면 HTML 이 오므로 **호출 전에** 막고 어디서 값을 얻는지 알린다."""
    with pytest.raises(LawError) as e:
        LawClient(oc="stub", throttle=0.0).body("eflaw", "286065")
    assert "efYd" in str(e.value) and "시행일자" in str(e.value)


def test_eflaw_body_normalizes_the_date_it_sends():
    """정규화된 '2026-11-20' 을 그대로 보내면 서버가 못 알아듣는다 — YYYYMMDD 로 되돌린다."""
    sent = {}

    class C(LawClient):
        def _call(self, url, params):
            sent.update(params)
            return b"<\xeb\xb2\x95\xeb\xa0\xb9/>"   # <법령/>

    C(oc="stub", throttle=0.0).body("eflaw", "286065", ef_date="2026-11-20")
    assert sent["efYd"] == "20261120"
    assert sent["MST"] == "286065"
